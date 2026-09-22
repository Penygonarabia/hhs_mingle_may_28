#!/usr/bin/env python3
"""
Generate a clean, high-impact A4 Landscape Word Document (.docx)
containing ONLY side-by-side screenshots:
Left: Management Presentation PDF Chart
Right: Staging Live Dashboard Screenshot (Sales Analysis with Salesman)

File name: Sales_Dashboard_Screenshots_Comparison.docx
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENTATION
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from PIL import Image

def set_cell_background(cell, hex_color):
    """Set background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=20, bottom=20, left=40, right=40):
    """Set cell internal margins (padding) in dxa."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    """Apply clean subtle borders to a table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def set_row_cant_split(row):
    """Prevent a table row from breaking across pages."""
    trPr = row._tr.get_or_add_trPr()
    cantSplit = parse_xml(f'<w:cantSplit {nsdecls("w")}/>')
    trPr.append(cantSplit)

def get_fitted_image_dimensions(img_path, max_w_in=5.35, max_h_in=6.20):
    """Calculate scaled width and height in Inches to fit within bounding box."""
    if not os.path.exists(img_path):
        return Inches(max_w_in), Inches(max_h_in)
    with Image.open(img_path) as im:
        w_px, h_px = im.size
    aspect = w_px / h_px
    calc_h = max_w_in / aspect
    if calc_h <= max_h_in:
        return Inches(max_w_in), Inches(calc_h)
    else:
        calc_w = max_h_in * aspect
        return Inches(calc_w), Inches(max_h_in)

def build_screenshots_document():
    doc = docx.Document()
    
    # Configure A4 Landscape with compact margins to maximize screenshot viewing area
    section = doc.sections[0]
    section.page_width = Inches(11.69)   # 297 mm
    section.page_height = Inches(8.27)   # 210 mm
    section.orientation = WD_ORIENTATION.LANDSCAPE
    section.left_margin = Inches(0.35)
    section.right_margin = Inches(0.35)
    section.top_margin = Inches(0.30)
    section.bottom_margin = Inches(0.30)

    # Set default normal style font and zero paragraph spacing
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(9)
    font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    doc.styles['Normal'].paragraph_format.space_before = Pt(0)
    doc.styles['Normal'].paragraph_format.space_after = Pt(0)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.0

    base_img_dir = '/Users/saravanan/Projects/hhs_cloud/docs/images/comparison'
    if not os.path.exists(base_img_dir) and os.path.exists('/Users/saravanan/Projects/cloud/docs/images/comparison'):
        base_img_dir = '/Users/saravanan/Projects/cloud/docs/images/comparison'
    pdf_img_dir = os.path.join(base_img_dir, 'pdf_pages')

    print("Building Screenshots-Only Comparison Document (A4 Landscape)...")

    # List of all comparison pages
    pages_data = [
        # Chart 1
        ("Chart 1", "Total Company Sales vs. Target & Last Year by Value & QTY", None, "pdf_page-01.png", "page_01.png"),
        # Chart 2
        ("Chart 2", "Sales Achievement and Growth Trends by Department & Region", None, "pdf_page-02.png", "page_02.png"),
        # Chart 3
        ("Chart 3", "Total Company Sales Quarterly Sales Progression", None, "pdf_page-03.png", "page_03.png"),
        # Chart 4 (Tabs: Q1, Q2, Q3)
        ("Chart 4", "Quarterly Sales Progression by Department", "Q1", "pdf_page-04.png", "page_04_tab_Q1.png"),
        ("Chart 4", "Quarterly Sales Progression by Department", "Q2", "pdf_page-04.png", "page_04_tab_Q2.png"),
        ("Chart 4", "Quarterly Sales Progression by Department", "Q3", "pdf_page-04.png", "page_04_tab_Q3.png"),
        # Chart 5 (Tabs: Q1, Q2, Q3)
        ("Chart 5", "Dealers Department Quarterly Sales Progression by Region", "Q1", "pdf_page-05.png", "page_05_tab_Q1.png"),
        ("Chart 5", "Dealers Department Quarterly Sales Progression by Region", "Q2", "pdf_page-05.png", "page_05_tab_Q2.png"),
        ("Chart 5", "Dealers Department Quarterly Sales Progression by Region", "Q3", "pdf_page-05.png", "page_05_tab_Q3.png"),
        # Chart 6
        ("Chart 6", "Total Company Sales vs. Target & Last Year by Group by Value", None, "pdf_page-06.png", "page_06.png"),
        # Chart 7
        ("Chart 7", "Product Groups — Top 10 by Value & QTY", None, "pdf_page-07.png", "page_07.png"),
        # Chart 8
        ("Chart 8", "Split Sales Distribution by Category (LY vs TY)", None, "pdf_page-08.png", "page_08.png"),
        # Chart 9
        ("Chart 9", "Main Categories Distribution by Value (LY vs TY)", None, "pdf_page-09.png", "page_09.png"),
        # Chart 10
        ("Chart 10", "Sales vs Target & Last Year by Department by Value & QTY", None, "pdf_page-10.png", "page_10.png"),
        # Chart 11
        ("Chart 11", "Channel Distribution by Value & QTY (LY vs TY)", None, "pdf_page-11.png", "page_11.png"),
        # Chart 12
        ("Chart 12", "Top Main Category — Sub-Categories (LCAC / RAC)", None, "pdf_page-12.png", "page_12.png"),
        # Chart 13 (Tabs: Dealers, Projects)
        ("Chart 13", "Top-3 Sub-Categories by Sales Type Group (YTD)", "Dealers", "pdf_page-13.png", "page_13_tab_Dealers.png"),
        ("Chart 13", "Top-3 Sub-Categories by Sales Type Group (YTD)", "Projects", "pdf_page-13.png", "page_13_tab_Projects.png"),
        # Chart 14 (Tabs: Dealers, Projects)
        ("Chart 14", "Regions by Sales Type Group (YTD)", "Dealers", "pdf_page-14.png", "page_14_tab_Dealers.png"),
        ("Chart 14", "Regions by Sales Type Group (YTD)", "Projects", "pdf_page-14.png", "page_14_tab_Projects.png"),
        # Chart 15 (Tabs: 4 Regions)
        ("Chart 15", "Regions — Main Categories (YTD)", "Western Region", "pdf_page-15.png", "page_15_tab_Western_Region.png"),
        ("Chart 15", "Regions — Main Categories (YTD)", "Riyadh Region", "pdf_page-15.png", "page_15_tab_Riyadh_Region.png"),
        ("Chart 15", "Regions — Main Categories (YTD)", "Eastern Region", "pdf_page-15.png", "page_15_tab_Eastern_Region.png"),
        ("Chart 15", "Regions — Main Categories (YTD)", "Qassim Region", "pdf_page-15.png", "page_15_tab_Qassim_Region.png"),
        # Chart 16 (Tabs: 4 Regions)
        ("Chart 16", "Regions — Sub-Categories (Split Models)", "Western Region", "pdf_page-16.png", "page_16_tab_Western_Region.png"),
        ("Chart 16", "Regions — Sub-Categories (Split Models)", "Riyadh Region", "pdf_page-16.png", "page_16_tab_Riyadh_Region.png"),
        ("Chart 16", "Regions — Sub-Categories (Split Models)", "Eastern Region", "pdf_page-16.png", "page_16_tab_Eastern_Region.png"),
        ("Chart 16", "Regions — Sub-Categories (Split Models)", "Qassim Region", "pdf_page-16.png", "page_16_tab_Qassim_Region.png"),
        # Chart 17
        ("Chart 17", "Top Sales Type Group — Product Groups (Category Mix)", None, "pdf_page-17.png", "page_17.png"),
    ]

    for p_idx, (chart_num, title, tab_name, pdf_file, dash_file) in enumerate(pages_data):
        if p_idx > 0:
            doc.add_page_break()

        # 1. Main Title Header Banner
        header_tbl = doc.add_table(rows=1, cols=1)
        header_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        header_tbl.autofit = False
        header_tbl.columns[0].width = Inches(10.99)
        h_cell = header_tbl.cell(0, 0)
        set_cell_background(h_cell, "0F172A")
        set_cell_margins(h_cell, top=30, bottom=30, left=80, right=80)
        set_row_cant_split(header_tbl.rows[0])

        hp = h_cell.paragraphs[0]
        hp.paragraph_format.space_before = Pt(0)
        hp.paragraph_format.space_after = Pt(0)
        tab_suffix = f"  —  [TAB: {tab_name.upper()}]" if tab_name else ""
        hrun = hp.add_run(f"{chart_num.upper()}: {title.upper()}{tab_suffix}")
        hrun.font.bold = True
        hrun.font.size = Pt(9.5)
        hrun.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        # 2. Side-by-Side Comparison Table (Headers + Images)
        comp_tbl = doc.add_table(rows=2, cols=2)
        comp_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        comp_tbl.autofit = False
        set_table_borders(comp_tbl, color="CBD5E1", sz="4")
        comp_tbl.columns[0].width = Inches(5.49)
        comp_tbl.columns[1].width = Inches(5.50)

        for row in comp_tbl.rows:
            set_row_cant_split(row)

        # Column Subheaders
        h0 = comp_tbl.cell(0, 0)
        h1 = comp_tbl.cell(0, 1)
        set_cell_background(h0, "1E293B")
        set_cell_background(h1, "1E3A8A")
        set_cell_margins(h0, top=20, bottom=20, left=40, right=40)
        set_cell_margins(h1, top=20, bottom=20, left=40, right=40)

        p_h0 = h0.paragraphs[0]
        p_h0.paragraph_format.space_before = Pt(0)
        p_h0.paragraph_format.space_after = Pt(0)
        p_h0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_h0 = p_h0.add_run("PDF Chart (Management Presentation Deck)")
        r_h0.font.bold = True
        r_h0.font.size = Pt(8.5)
        r_h0.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        p_h1 = h1.paragraphs[0]
        p_h1.paragraph_format.space_before = Pt(0)
        p_h1.paragraph_format.space_after = Pt(0)
        p_h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        dash_label = f"Staging Dashboard (Sales Analysis with Salesman)" + (f" — Tab: {tab_name}" if tab_name else "")
        r_h1 = p_h1.add_run(dash_label)
        r_h1.font.bold = True
        r_h1.font.size = Pt(8.5)
        r_h1.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        # Image Cells (Maximizing screenshot display area)
        c_left = comp_tbl.cell(1, 0)
        c_right = comp_tbl.cell(1, 1)
        set_cell_background(c_left, "FAFAFA")
        set_cell_background(c_right, "FAFAFA")
        set_cell_margins(c_left, top=20, bottom=20, left=20, right=20)
        set_cell_margins(c_right, top=20, bottom=20, left=20, right=20)

        # Left Image: PDF Page
        p_left = c_left.paragraphs[0]
        p_left.paragraph_format.space_before = Pt(0)
        p_left.paragraph_format.space_after = Pt(0)
        p_left.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pdf_path = os.path.join(pdf_img_dir, pdf_file)
        w_l, h_l = get_fitted_image_dimensions(pdf_path, max_w_in=5.35, max_h_in=6.50)
        if os.path.exists(pdf_path):
            p_left.add_run().add_picture(pdf_path, width=w_l, height=h_l)

        # Right Image: Staging Dashboard Page / Tab
        p_right = c_right.paragraphs[0]
        p_right.paragraph_format.space_before = Pt(0)
        p_right.paragraph_format.space_after = Pt(0)
        p_right.alignment = WD_ALIGN_PARAGRAPH.CENTER
        dash_path = os.path.join(base_img_dir, dash_file)
        w_r, h_r = get_fitted_image_dimensions(dash_path, max_w_in=5.35, max_h_in=6.50)
        if os.path.exists(dash_path):
            p_right.add_run().add_picture(dash_path, width=w_r, height=h_r)

    out_file_primary = '/Users/saravanan/Projects/hhs_cloud/docs/Sales_Dashboard_Screenshots_Comparison.docx'
    doc.save(out_file_primary)
    print(f"✅ Screenshots-only comparison document created: {out_file_primary}")

    out_file_cloud = '/Users/saravanan/Projects/cloud/docs/Sales_Dashboard_Screenshots_Comparison.docx'
    if os.path.exists('/Users/saravanan/Projects/cloud/docs'):
        doc.save(out_file_cloud)
        print(f"✅ Document also copied to: {out_file_cloud}")

if __name__ == '__main__':
    build_screenshots_document()
