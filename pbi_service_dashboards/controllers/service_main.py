# -*- coding: utf-8 -*-
"""Backs every "PBI Dashboards > Service Dashboards" leaf — a single
generic controller for all 15 dashboards (see service_config.py's
``BOARDS`` registry), each identified by a ``board`` key. Reads live from
``project_task``/``machine_repair_support``/``mail_message`` directly
(see service_sql.py) — no ``ks_dashboard_ninja`` or ``dbmodel.*.ct``
dependency.

Access is gated the same way as every other pbi_dashboards menu. When
user_menu_rights is installed (an optional overlay, not a dependency) the
menu is hidden until an admin grants it there; otherwise ordinary Odoo
groups decide. Either way it is
enforced both on the menu (existing mechanism) and here on the JSON data
routes themselves — same double-gate pattern as sales_kpi_main.py.
"""
from odoo import http
from odoo.http import request

from . import service_sql
from .service_config import BOARDS, resolve_scope_info, resolve_item_info
from odoo.addons.pbi_dashboards.controllers.board_engine import PbiDashboardBoardEngineMixin  # noqa: F401


# The Formula & Details modal lets the user narrow the table on screen
# (search box + column sort) before hitting Download. The download is a
# plain browser navigation to a separate request, so the server re-runs
# the query and would otherwise hand back the *unfiltered* list — a file
# that does not match the table the user was looking at. These helpers
# replay the client's search/sort on the server side, over the same rows,
# with the same rules as service_dashboard.js's filteredDetailRecords (the
# entity filter is not here: it is pushed down into SQL by
# run_chart_detail's entity_id argument).
#
# Which columns exist varies per chart, so the searchable set is derived
# from the payload rather than hard-coded.
_BASE_RECORD_COLS = ("name", "entity_name", "work_center", "region", "status", "value_formatted")
# columns the JS sorts ascending by default on first click
_TEXT_SORT_COLS = ("name", "entity_name", "work_center", "region", "status")


def _region_label(region):
    """The region filter as a reader sees it. The dropdown's "no restriction"
    value is the literal string 'all', which is truthy — printing it raw put
    "Region,all" at the top of every unfiltered export."""
    if not region or region == "all":
        return "All Regions"
    return region


def _franchise_label(franchise):
    """The franchise filter as a reader sees it — same 'all' guard as
    _region_label above, so an unfiltered export does not print
    "Franchise,all" in its header block."""
    if not franchise or franchise == "all":
        return "All Franchises"
    return franchise


def _searchable_cols(columns):
    return _BASE_RECORD_COLS + tuple(c["key"] for c in (columns or []))


def _apply_view_filters(records, columns=None, search=None, sort_by=None, sort_asc=None):
    """Return records narrowed and ordered exactly as the modal shows them."""
    rows = list(records or [])

    query = (search or "").strip().lower()
    if query:
        cols = _searchable_cols(columns)
        rows = [r for r in rows
                if any(query in str(r.get(col) or "").lower() for col in cols)]

    if sort_by:
        if sort_asc is None:
            asc = sort_by in _TEXT_SORT_COLS
        elif isinstance(sort_asc, str):
            asc = sort_asc.lower() in ("1", "true", "yes")
        else:
            asc = bool(sort_asc)

        def key(rec):
            val = rec.get(sort_by)
            if isinstance(val, (int, float)):
                return (0, val, "")
            return (1, 0.0, str(val).lower())

        # A record with no value sinks to the bottom in BOTH directions —
        # reversing it into first place would bury the rows the reader
        # opened the table to see. Sorted apart from the rest so the
        # reverse flag cannot reach it.
        blanks = [r for r in rows if r.get(sort_by) is None]
        rest = [r for r in rows if r.get(sort_by) is not None]
        # numbers and strings never mix within one column, so the leading
        # type tag only keeps the tuples comparable — it never reorders.
        rest.sort(key=key, reverse=not asc)
        rows = rest + blanks

    return rows


def _fmt_hours(hours):
    """0.0-based float hours -> "H:MM", same rounding as the JS/SQL side."""
    neg = hours < 0
    hours = abs(hours)
    h_int = int(hours)
    m_int = int(round((hours - h_int) * 60))
    if m_int >= 60:
        h_int += 1
        m_int = 0
    return "%s%d:%02d" % ("-" if neg else "", h_int, m_int)


def _column_totals(records, columns):
    """Sums for the extra columns whose DetailColumn asked for one — the
    scheduling formula's "Number of New Jobs" divisor is the one that does.
    Same figure the modal's footer shows (service_dashboard.js's
    detailColumnTotals), over the same rows, so a downloaded file carries
    the divisor rather than making the reader add a column up by hand.

    Read off the *_raw values run_chart_detail sends beside each cell: the
    displayed one is already a formatted string."""
    out = {}
    for col in columns or []:
        if not col.get("total"):
            continue
        total = sum(float(rec.get(col["key"] + "_raw") or 0) for rec in records)
        out[col["key"]] = int(total) if total == int(total) else round(total, 2)
    return out


def _view_totals(records, summary):
    """Totals for the rows actually being exported, not for the whole period.

    Mirrors the aggregate the chart itself applies: a sum chart totals, an
    average chart divides by the rows that actually have a value (a job
    card whose interval never completed is listed but not counted), a
    count chart just counts, and utilization divides the total by the
    period-scaled 176-hour figure the engine already computed.
    """
    kind = summary.get("value_kind", "hours")
    agg = summary.get("agg", "sum")
    counted = [r for r in records if r.get("counted") and r.get("value") is not None]
    total = sum(float(r.get("value") or 0.0) for r in counted)
    avg = (total / len(counted)) if counted else 0.0

    def fmt(v):
        return _fmt_hours(v) if kind == "hours" else f"{v:,.2f}"

    totals = {
        "agg": agg,
        "count": len(records),
        "counted": len(records) if agg == "count" else len(counted),
        "total": round(total, 2),
        "total_formatted": fmt(total),
        "avg": round(avg, 2),
        "avg_formatted": fmt(avg),
    }
    if agg == "utilization":
        divisor = float(summary.get("divisor") or 0.0)
        totals["divisor"] = divisor
        totals["divisor_label"] = summary.get("divisor_label", "")
        totals["utilization_pct"] = round((total / divisor * 100.0) if divisor else 0.0, 2)
    return totals


class _ServiceMixin(PbiDashboardBoardEngineMixin):
    """Thin override layer: supplies service-specific item info strings from
    resolve_item_info (service_config.py). All other mixin methods are
    inherited unchanged from PbiDashboardBoardEngineMixin."""

    def _resolve_item_info(self, board_cfg, item_cfg):
        return resolve_item_info(board_cfg, item_cfg)


class PbiServiceDashboardController(_ServiceMixin, http.Controller):

    @http.route('/pbi_dashboards/service/board', type='json', auth='user')
    def board_data(self, board, dateFilter=None, customStart=None, customEnd=None, region=None,
                   franchise=None):
        if not self._has_access(board):
            return {'error': 'You do not have access to this dashboard.'}
        board_cfg = BOARDS.get(board)
        if not board_cfg:
            return {'error': 'Unknown dashboard.'}
        try:
            date_from, date_to, resolved_filter, label = self._resolve_date_range(dateFilter, customStart, customEnd)
            effective_board = self._effective_board(board_cfg, region, franchise)
            uid = request.env.uid
            period_months = self._period_months(date_from, date_to)
            items = [self._item_payload(effective_board, item_cfg, date_from, date_to, uid, period_months)
                     for item_cfg in board_cfg.items]
            result = {
                "title": board_cfg.title,
                "scopeInfo": resolve_scope_info(board_cfg),
                "period": {"dateFilter": resolved_filter, "label": label},
                "regionFilterable": board_cfg.region_filterable,
                # Offered on every board, unlike region: the franchise lives
                # on the job card itself, so there is no board it cannot
                # narrow (see board_engine._effective_board).
                "franchiseOptions": self._franchise_options(request.env),
                "items": items,
            }
            if board_cfg.region_filterable:
                result["regionOptions"] = self._region_options(request.env)
            return result
        except Exception as e:
            return {'error': str(e)}

    def _chart_detail(self, board_cfg, item_cfg, dateFilter, customStart, customEnd, region, entity_id,
                      franchise=None):
        """The Formula & Details payload for one chart, plus the context
        strings the modal and the exports print above the table."""
        fn = getattr(service_sql, 'run_chart_detail', None)
        if not fn:
            raise ValueError('Formula & Details engine is not available. '
                             'Please update the pbi_dashboards module on this server.')
        if not getattr(item_cfg, 'detail', None):
            raise ValueError(f'"{item_cfg.name}" has no formula breakdown configured.')
        date_from, date_to, _resolved_filter, label = self._resolve_date_range(dateFilter, customStart, customEnd)
        effective_board = self._effective_board(board_cfg, region, franchise)
        data = fn(request.env, request.env.uid, effective_board, item_cfg, date_from, date_to,
                  entity_id=entity_id, period_months=self._period_months(date_from, date_to))
        data["periodLabel"] = label
        data["boardTitle"] = board_cfg.title
        data["chartName"] = item_cfg.name
        data["region"] = _region_label(region)
        data["franchise"] = _franchise_label(franchise)
        return data

    @http.route('/pbi_dashboards/service/chart', type='json', auth='user')
    def chart_data(self, board, item, drillPath=None, dateFilter=None, customStart=None, customEnd=None,
                   region=None, franchise=None, details=False, entityId=None, entityIds=None,
                   technicianId=None, **kw):
        if not self._has_access(board):
            return {'error': 'You do not have access to this dashboard.'}
        board_cfg = BOARDS.get(board)
        if not board_cfg:
            return {'error': 'Unknown dashboard.'}
        item_cfg = next((i for i in board_cfg.items if i.key == item), None)
        if not item_cfg:
            return {'error': 'Unknown chart.'}
        try:
            date_from, date_to, _resolved_filter, label = self._resolve_date_range(dateFilter, customStart, customEnd)
            effective_board = self._effective_board(board_cfg, region, franchise)
            uid = request.env.uid

            if details:
                # entityIds (plural) is what the multi-select picker sends;
                # entityId is the single-select parameter that preceded it and
                # technicianId the pre-generalisation name before that — both
                # still accepted so an older cached bundle keeps working.
                return self._chart_detail(board_cfg, item_cfg, dateFilter, customStart, customEnd,
                                          region, entityIds or entityId or technicianId,
                                          franchise)

            drill_path = drillPath if isinstance(drillPath, list) else []
            period_months = self._period_months(date_from, date_to)
            breakdown = service_sql.run_breakdown(request.env, uid, effective_board, item_cfg, date_from, date_to,
                                                  drill_path, period_months)
            if breakdown is None:
                model, domain = service_sql.run_terminal_domain(request.env, uid, effective_board, item_cfg, drill_path, date_from, date_to)
                return {'terminal': True, 'model': model, 'domain': domain, 'name': item_cfg.name,
                        'listViewId': self._flat_list_view_id(request.env, model)}
            return {'terminal': False, 'breakdown': breakdown, 'level': len(drill_path)}
        except Exception as e:
            return {'error': str(e)}


    # -----------------------------------------------------------------
    # Formula & Details exports. Both formats show the same rows the modal
    # is showing (see _apply_view_filters) under the same formula header,
    # so a downloaded file explains itself and reconciles with the chart.
    # -----------------------------------------------------------------
    def _detail_export_data(self, board, item, dateFilter, customStart, customEnd,
                            region, entityId, search, sortBy, sortAsc, franchise=None):
        board_cfg = BOARDS[board]
        item_cfg = next(i for i in board_cfg.items if i.key == item)
        data = self._chart_detail(board_cfg, item_cfg, dateFilter, customStart, customEnd,
                                  region, entityId, franchise)
        columns = data.get("columns", [])
        records = _apply_view_filters(data.get("records", []), columns, search, sortBy, sortAsc)
        data["records"] = records
        data["totals"] = _view_totals(records, data.get("summary", {}))
        data["columnTotals"] = _column_totals(records, columns)
        return data

    def _detail_export_meta(self, data, region, search, franchise=None):
        """The self-explaining header block both formats print above the
        table: what the chart is, over what period, under what formula."""
        summary = data.get("summary", {})
        formula = data.get("formula", {})
        totals = data.get("totals", {})
        agg = totals.get("agg", "sum")

        rows = [
            ("Dashboard", data.get("boardTitle", "")),
            ("Chart", data.get("chartName", "")),
            ("Period", data.get("periodLabel", "")),
            ("Region", _region_label(region)),
            ("Franchise", _franchise_label(franchise)),
            ("Formula Used", formula.get("expression", "")),
        ]
        rows += [(label, definition) for label, definition in
                 ((t.get("label", ""), t.get("definition", "")) for t in formula.get("terms", []))]
        if formula.get("scope"):
            rows.append(("Scope", formula["scope"]))
        if search:
            rows.append(("Search Filter", search))

        record_label = summary.get("record_label", "Records")
        rows.append((f"Total {record_label}", str(totals.get("count", 0))))
        if agg != "count":
            value_label = summary.get("value_label", "Value")
            if agg == "avg":
                rows.append(("Counted (with a value)", str(totals.get("counted", 0))))
                rows.append((f"Average {value_label}", totals.get("avg_formatted", "-")))
            rows.append((f"Total {value_label}", totals.get("total_formatted", "-")))
        if agg == "utilization":
            rows.append(("Divisor", totals.get("divisor_label", "")))
            rows.append(("Utilization", f"{totals.get('utilization_pct', 0.0)}%"))
        return rows

    def _detail_export_headers(self, data):
        summary = data.get("summary", {})
        headers = ["#", "Job Card #", summary.get("entity_label", "Technician"),
                   "Work Center", "Region", "Status"]
        headers += [c["label"] for c in data.get("columns", [])]
        if data.get("valueColumn", {}).get("shown"):
            headers.append(summary.get("value_label", "Value"))
            headers.append(summary.get("value_label", "Value") + " (Decimal)")
            headers.append("Counted")
        return headers

    def _detail_export_row(self, data, index, rec):
        row = [index, rec.get("name"), rec.get("entity_name"), rec.get("work_center"),
               rec.get("region"), rec.get("status")]
        row += [rec.get(c["key"]) for c in data.get("columns", [])]
        if data.get("valueColumn", {}).get("shown"):
            row.append(rec.get("value_formatted"))
            row.append(rec.get("value"))
            row.append("Yes" if rec.get("counted") else "No")
        return row

    def _detail_export_filename(self, data, ext):
        parts = [data.get("chartName", "Chart Details"), data.get("periodLabel", "")]
        raw = "_".join(p for p in parts if p)
        safe = "".join(c for c in raw if c.isalnum() or c in (" ", "_", "-")).strip().replace(" ", "_")
        return f"{safe or 'Chart_Details'}.{ext}"

    @http.route('/pbi_dashboards/service/detail_export_xlsx', type='http', auth='user')
    def detail_export_xlsx(self, board, item, dateFilter=None, customStart=None, customEnd=None,
                           region=None, franchise=None, entityId=None, entityIds=None,
                           search=None, sortBy=None, sortAsc=None, **kw):
        if not self._has_access(board):
            return request.not_found()
        board_cfg = BOARDS.get(board)
        if not board_cfg or not any(i.key == item for i in board_cfg.items):
            return request.not_found()

        try:
            from datetime import datetime
            import io
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter

            data = self._detail_export_data(board, item, dateFilter, customStart, customEnd,
                                            region, entityIds or entityId, search, sortBy, sortAsc,
                                            franchise)

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Formula & Details"
            ws.views.sheetView[0].showGridLines = True

            font_title = Font(name="Calibri", size=15, bold=True, color="1F4E79")
            font_sub = Font(name="Calibri", size=10, italic=True, color="595959")
            font_meta_hdr = Font(name="Calibri", size=10, bold=True, color="1F4E79")
            font_hdr = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            font_data = Font(name="Calibri", size=10)
            font_bold = Font(name="Calibri", size=10, bold=True)
            font_muted = Font(name="Calibri", size=10, color="9CA3AF", italic=True)
            fill_hdr = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
            fill_zebra = PatternFill(start_color="F2F7FA", end_color="F2F7FA", fill_type="solid")
            fill_box = PatternFill(start_color="EBF1F5", end_color="EBF1F5", fill_type="solid")
            fill_total = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
            thin_border = Border(
                left=Side(style='thin', color='D9D9D9'),
                right=Side(style='thin', color='D9D9D9'),
                top=Side(style='thin', color='D9D9D9'),
                bottom=Side(style='thin', color='D9D9D9'),
            )
            total_border = Border(
                top=Side(style='thin', color='1F4E79'),
                bottom=Side(style='double', color='1F4E79'),
            )

            headers = self._detail_export_headers(data)
            n_cols = len(headers)

            ws.cell(row=1, column=1,
                    value=f"{data.get('boardTitle', '')} - {data.get('chartName', '')} - Calculation Breakdown").font = font_title
            ws.cell(row=2, column=1,
                    value=f"Exported: {datetime.now().strftime('%Y-%m-%d %H:%M')}").font = font_sub

            meta = self._detail_export_meta(data, region, search, franchise)
            row_idx = 4
            for label, value in meta:
                ws.cell(row=row_idx, column=1, value=label).font = font_meta_hdr
                cell = ws.cell(row=row_idx, column=2, value=value)
                cell.font = font_bold if label.startswith(("Total", "Average", "Utilization")) else font_data
                cell.alignment = Alignment(wrap_text=False)
                for c in range(1, max(n_cols, 2) + 1):
                    ws.cell(row=row_idx, column=c).fill = fill_box
                row_idx += 1

            start_row = row_idx + 1
            for col_idx, h in enumerate(headers, 1):
                cell = ws.cell(row=start_row, column=col_idx, value=h)
                cell.font = font_hdr
                cell.fill = fill_hdr
                cell.alignment = Alignment(horizontal="center" if col_idx > 5 or col_idx == 1 else "left",
                                           vertical="center")

            records = data.get("records", [])
            row_idx = start_row + 1
            for i, rec in enumerate(records, 1):
                for col_idx, val in enumerate(self._detail_export_row(data, i, rec), 1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=val)
                    # A row the average skipped is greyed, not hidden — it
                    # is why the count and the denominator differ.
                    cell.font = font_data if rec.get("counted") else font_muted
                    if i % 2 == 0:
                        cell.fill = fill_zebra
                    cell.border = thin_border
                    if col_idx > 5:
                        cell.alignment = Alignment(horizontal="center")
                row_idx += 1

            totals = data.get("totals", {})
            summary = data.get("summary", {})
            ws.cell(row=row_idx, column=1, value="Total").font = font_bold
            ws.cell(row=row_idx, column=2,
                    value=f"{totals.get('count', 0)} {summary.get('record_label', 'Records')}").font = font_bold
            # ...and under each extra column that carries a total of its own
            # (the 6 leading columns are #/Job Card/entity/Work Center/Region/Status)
            col_totals = data.get("columnTotals", {})
            for i, c in enumerate(data.get("columns", [])):
                if c["key"] in col_totals:
                    cell = ws.cell(row=row_idx, column=7 + i, value=col_totals[c["key"]])
                    cell.font = font_bold
                    cell.alignment = Alignment(horizontal="center")
            if data.get("valueColumn", {}).get("shown"):
                value_col = 6 + len(data.get("columns", [])) + 1
                ws.cell(row=row_idx, column=value_col, value=totals.get("total_formatted")).font = font_bold
                ws.cell(row=row_idx, column=value_col + 1, value=totals.get("total")).font = font_bold
            for c in range(1, n_cols + 1):
                cell = ws.cell(row=row_idx, column=c)
                cell.fill = fill_total
                cell.border = total_border

            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                # the formula/definition text in column B would otherwise
                # blow the sheet out to several screens wide
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 60)

            output = io.BytesIO()
            wb.save(output)
            output.seek(0)

            filename = self._detail_export_filename(data, "xlsx")
            resp_headers = [
                ('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
                ('Content-Disposition', f'attachment; filename="{filename}"'),
            ]
            return request.make_response(output.getvalue(), headers=resp_headers)
        except Exception as e:
            return request.make_response(f"Export failed: {str(e)}", status=500)

    @http.route('/pbi_dashboards/service/detail_export_csv', type='http', auth='user')
    def detail_export_csv(self, board, item, dateFilter=None, customStart=None, customEnd=None,
                          region=None, franchise=None, entityId=None, entityIds=None,
                          search=None, sortBy=None, sortAsc=None, **kw):
        """Same rows as the xlsx export, as plain CSV — for users who want
        to pivot the data rather than read the formatted sheet, and for
        servers where openpyxl is not installed (the xlsx route needs it,
        this one does not)."""
        if not self._has_access(board):
            return request.not_found()
        board_cfg = BOARDS.get(board)
        if not board_cfg or not any(i.key == item for i in board_cfg.items):
            return request.not_found()

        try:
            import csv
            import io
            from datetime import datetime

            data = self._detail_export_data(board, item, dateFilter, customStart, customEnd,
                                            region, entityIds or entityId, search, sortBy, sortAsc,
                                            franchise)

            buf = io.StringIO()
            writer = csv.writer(buf, lineterminator="\r\n")
            writer.writerow([f"{data.get('boardTitle', '')} - {data.get('chartName', '')} - Calculation Breakdown"])
            writer.writerow(["Exported", datetime.now().strftime("%Y-%m-%d %H:%M")])
            for label, value in self._detail_export_meta(data, region, search, franchise):
                writer.writerow([label, value])
            writer.writerow([])

            writer.writerow(self._detail_export_headers(data))
            for i, rec in enumerate(data.get("records", []), 1):
                writer.writerow(self._detail_export_row(data, i, rec))

            totals = data.get("totals", {})
            summary = data.get("summary", {})
            col_totals = data.get("columnTotals", {})
            total_row = ["Total", f"{totals.get('count', 0)} {summary.get('record_label', 'Records')}"]
            total_row += [""] * 4
            total_row += [col_totals.get(c["key"], "") for c in data.get("columns", [])]
            if data.get("valueColumn", {}).get("shown"):
                total_row += [totals.get("total_formatted"), totals.get("total"), ""]
            writer.writerow(total_row)

            filename = self._detail_export_filename(data, "csv")
            # UTF-8 BOM: without it Excel on Windows reads the Arabic
            # technician/work-centre names as mojibake.
            payload = "﻿" + buf.getvalue()
            resp_headers = [
                ('Content-Type', 'text/csv; charset=utf-8'),
                ('Content-Disposition', f'attachment; filename="{filename}"'),
            ]
            return request.make_response(payload.encode("utf-8"), headers=resp_headers)
        except Exception as e:
            return request.make_response(f"Export failed: {str(e)}", status=500)
